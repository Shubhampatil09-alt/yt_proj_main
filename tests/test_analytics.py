import csv
import io
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from app.database import JobHistory


def iso(day: int, hour: int = 0) -> str:
    return datetime(2026, 1, day, hour, tzinfo=timezone.utc).isoformat(timespec="seconds")


def add_run(store, session, run_id, worker_id, status, started, duration=10, progress=100, error=None):
    store.begin_run(run_id, session, worker_id, started)
    store.record_transition(run_id, worker_id, "Running", started)
    store.record_transition(run_id, worker_id, "Running", started)
    store.finalize_run(
        run_id,
        ended_at=(datetime.fromisoformat(started) + timedelta(seconds=duration)).isoformat(timespec="seconds"),
        execution_duration=duration,
        final_status=status,
        playback_duration=12,
        playback_progress=progress,
        error_message=error,
        browser_process_info={"pids": [100 + worker_id], "memory_mb": 64.5},
        stop_reason="safety:ram" if status == "safety_stopped" else None,
    )


def make_store(tmp_path):
    store = JobHistory(tmp_path / "analytics.db")
    store.initialize()
    store.start_session("session-a", iso(1))
    return store


def test_execution_history_and_transitions_are_created_once(tmp_path):
    store = make_store(tmp_path)
    add_run(store, "session-a", "run-a", 7, "completed", iso(2), duration=4.25)
    duplicate_finalize = store.finalize_run(
        "run-a", ended_at=iso(3), execution_duration=99, final_status="failed",
        playback_duration=0, playback_progress=0, error_message="duplicate",
        browser_process_info={}, stop_reason=None,
    )

    run = store.get_run("run-a")

    assert duplicate_finalize is False
    assert run["final_status"] == "completed"
    assert run["execution_duration"] == 4.25
    assert [transition["state"] for transition in run["transitions"]] == ["Starting", "Running"]


def test_analytics_calculations_use_stored_records(tmp_path):
    store = make_store(tmp_path)
    add_run(store, "session-a", "one", 1, "completed", iso(2), duration=10)
    add_run(store, "session-a", "two", 1, "failed", iso(3), duration=20, progress=40, error="codec")
    add_run(store, "session-a", "three", 2, "safety_stopped", iso(4), duration=30, progress=60)
    store.insert_resource_samples([
        {"session_id":"session-a","timestamp":iso(2),"cpu_percent":25,"ram_percent":50,"available_ram_gb":4,"worker_ram_mb":80,"active_worker_count":1,"worker_memory":{"1":80}},
        {"session_id":"session-a","timestamp":iso(3),"cpu_percent":92,"ram_percent":88,"available_ram_gb":1,"worker_ram_mb":190,"active_worker_count":2,"worker_memory":{"1":90,"2":100}},
    ])

    summary = store.analytics_summary({})

    assert summary["total_runs"] == 3
    assert summary["successful_runs"] == 1
    assert summary["failed_runs"] == 1
    assert summary["success_percentage"] == 33.33
    assert summary["average_duration"] == 20
    assert summary["maximum_duration"] == 30
    assert summary["peak_cpu"] == 92
    assert summary["peak_ram"] == 88
    assert summary["peak_worker_ram"] == 190
    assert summary["safety_shutdowns"] == 1
    assert summary["worker_failure_frequency"] == [{"worker_id": 1, "failures": 1}]


def test_resource_samples_are_buffered_then_flushed(client):
    client.get("/api/metrics")

    assert client.app.state.analytics.buffered_count == 1
    response = client.get("/api/analytics/resources")

    assert response.status_code == 200
    assert len(response.json()["items"]) == 1
    assert response.json()["items"][0]["active_worker_count"] == 1


def test_date_status_filters_and_pagination(client):
    store = client.app.state.dashboard.history
    session = client.app.state.session_id
    for number in range(30):
        status = "failed" if number % 2 else "completed"
        add_run(store, session, f"run-{number:02}", number % 3 + 1, status, iso(2 + number // 20, number % 20), error="x" if status == "failed" else None)

    response = client.get(
        "/api/analytics/runs",
        params={"status":"failed","start":iso(2),"end":iso(3,23),"page":2,"page_size":5,"sort_by":"worker_id","sort_order":"asc"},
    )
    invalid = client.get("/api/analytics/runs", params={"status":"unknown"})

    payload = response.json()
    assert response.status_code == 200
    assert payload["total"] == 15
    assert payload["page"] == 2
    assert len(payload["items"]) == 5
    assert all(item["final_status"] == "failed" for item in payload["items"])
    assert invalid.status_code == 422


def test_csv_and_excel_exports_are_valid(client):
    store = client.app.state.dashboard.history
    session = client.app.state.session_id
    add_run(store, session, "export-run", 2, "failed", iso(2), error="playback failed")
    client.get("/api/metrics")

    csv_response = client.get("/api/analytics/export", params={"format":"csv"})
    xlsx_response = client.get("/api/analytics/export", params={"format":"xlsx"})
    rows = list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    workbook = load_workbook(io.BytesIO(xlsx_response.content))

    assert csv_response.status_code == 200
    assert rows[0]["run_id"] == "export-run"
    assert xlsx_response.status_code == 200
    assert workbook.sheetnames == ["Summary", "Test Runs", "Resource Samples", "Errors"]
    assert workbook["Errors"].max_row == 2


def test_retention_never_deletes_active_session_data(tmp_path):
    store = make_store(tmp_path)
    store.start_session("old-session", "2020-01-01T00:00:00+00:00")
    add_run(store, "old-session", "expired", 1, "completed", "2020-01-01T00:00:00+00:00")
    add_run(store, "session-a", "active-old", 2, "completed", "2020-01-01T00:00:00+00:00")
    old_sample = {"timestamp":"2020-01-01T00:00:00+00:00","cpu_percent":1,"ram_percent":2,"available_ram_gb":7,"worker_ram_mb":0,"active_worker_count":0,"worker_memory":{}}
    store.insert_resource_samples([{**old_sample,"session_id":"old-session"},{**old_sample,"session_id":"session-a"}])

    removed = store.cleanup_retention("session-a", resource_days=7, run_days=90, now=datetime(2026,1,10,tzinfo=timezone.utc))

    assert removed == {"resource_samples": 1, "test_runs": 1}
    assert store.get_run("expired") is None
    assert store.get_run("active-old") is not None
    assert len(store.list_resources({"session_id":"session-a"})) == 1


def test_sqlite_concurrent_writes_are_serialized(tmp_path):
    store = make_store(tmp_path)

    def write(number):
        add_run(store, "session-a", f"parallel-{number}", number, "completed", iso(2), duration=number)

    with ThreadPoolExecutor(max_workers=8) as executor:
        list(executor.map(write, range(1, 41)))

    assert store.analytics_summary({})["total_runs"] == 40

