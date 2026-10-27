from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .database import JobHistory


RUN_HEADERS = [
    "run_id", "session_id", "worker_id", "bot_id", "bot_name", "started_at", "ended_at",
    "execution_duration", "final_status", "playback_duration",
    "playback_progress", "error_message", "browser_process_info", "stop_reason",
]


class ReportExporter:
    def __init__(self, store: JobHistory) -> None:
        self.store = store

    def csv(self, filters: dict[str, Any]) -> bytes:
        runs = self.store.list_runs(filters, page=1, page_size=1_000_000)["items"]
        output = io.StringIO(newline="")
        writer = csv.DictWriter(output, fieldnames=RUN_HEADERS, extrasaction="ignore")
        writer.writeheader()
        for run in runs:
            row = dict(run)
            row["browser_process_info"] = json.dumps(row["browser_process_info"])
            writer.writerow(row)
        return output.getvalue().encode("utf-8-sig")

    def xlsx(self, filters: dict[str, Any]) -> bytes:
        summary = self.store.analytics_summary(filters)
        runs = self.store.list_runs(filters, page=1, page_size=1_000_000)["items"]
        resources = self.store.list_resources(filters, limit=1_000_000)
        errors = self.store.list_errors(filters, limit=1_000_000)

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        self._write_summary(summary_sheet, summary)
        self._write_records(workbook.create_sheet("Test Runs"), runs, RUN_HEADERS)
        self._write_records(
            workbook.create_sheet("Resource Samples"),
            resources,
            ["session_id", "timestamp", "cpu_percent", "ram_percent", "available_ram_gb",
             "worker_ram_mb", "chrome_memory_mb", "chromedriver_memory_mb",
             "youtube_memory_mb", "backend_rss_mb", "active_worker_count", "worker_memory"],
        )
        self._write_records(
            workbook.create_sheet("Errors"), errors,
            ["run_id", "session_id", "worker_id", "bot_id", "bot_name",
             "started_at", "ended_at", "error_message"],
        )
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    @staticmethod
    def _write_summary(sheet, summary: dict[str, Any]) -> None:
        sheet.append(["Metric", "Value"])
        for key, value in summary.items():
            if isinstance(value, (dict, list)):
                value = json.dumps(value)
            sheet.append([key.replace("_", " ").title(), value])
        ReportExporter._style_sheet(sheet)
        sheet.column_dimensions["A"].width = 30
        sheet.column_dimensions["B"].width = 24

    @staticmethod
    def _write_records(sheet, records: list[dict[str, Any]], headers: list[str]) -> None:
        sheet.append([header.replace("_", " ").title() for header in headers])
        for record in records:
            row = []
            for header in headers:
                value = record.get(header)
                if isinstance(value, (dict, list)):
                    value = json.dumps(value)
                elif header in {"started_at", "ended_at", "timestamp"} and value:
                    try:
                        value = datetime.fromisoformat(str(value).replace("Z", "+00:00")).replace(tzinfo=None)
                    except ValueError:
                        pass
                row.append(value)
            sheet.append(row)
        ReportExporter._style_sheet(sheet)
        for row in sheet.iter_rows(min_row=2):
            for cell, header in zip(row, headers):
                if header in {"started_at", "ended_at", "timestamp"} and isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif header in {"execution_duration", "playback_duration"}:
                    cell.number_format = "0.00"
                elif header in {"playback_progress", "cpu_percent", "ram_percent"}:
                    cell.number_format = "0.0"

    @staticmethod
    def _style_sheet(sheet) -> None:
        header_fill = PatternFill("solid", fgColor="172033")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
            sheet.column_dimensions[get_column_letter(column)].width = min(48, max(12, max(map(len, values)) + 2))
